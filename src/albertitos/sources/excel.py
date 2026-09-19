"""El Excel de Alberto, limpio: sólo Proveedores, Pedidos_2026 y la norma. El resto son avisos.

Tolerante a la FORMA, estricto con el CONTENIDO: las columnas se buscan por cabecera normalizada
(sin tildes, sin mayúsculas, sin separadores), no por posición, y las hojas por nombre normalizado.
Lo que no se entiende va a `MasterSnapshot.avisos_calidad` con una frase que se pueda leer en voz alta.
Sólo se niega a cargar si falta la columna clave o el importe de los pedidos: un maestro medio vacío
escalaría las 500 facturas sin decir por qué, y eso es peor que un error con nombre.
"""

from __future__ import annotations

import re
from contextlib import closing
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import openpyxl

from albertitos.core.contracts import MasterSnapshot, Pedido, Proveedor
from albertitos.core.hashing import hash_canonico
from albertitos.formatos import (
    CENT,
    normalizar_iban,
    normalizar_nif,
    normalizar_pedido,
    parse_fecha_es,
    parse_importe_es,
    sin_tildes,
)

HOJAS_NOTAS = (
    "notas_alberto",
    "pendiente_revisar",
    "Pedidos_2025_OLD",
    "NO_TOCAR",
    "backup_marzo",
    "MACROS_ROTAS",
)
HOJA_NORMA_VIGENTE = "Norma_Pagos_v3"

#: Nombres aceptados para cada hoja de datos (normalizados; el primero es el canónico).
ALIAS_HOJAS: dict[str, tuple[str, ...]] = {
    "Proveedores": ("proveedores", "proveedor", "maestroproveedores", "suppliers"),
    "Pedidos_2026": ("pedidos2026", "pedidos", "pedido", "ordenes", "ordenescompra", "orders"),
}

#: Cabeceras aceptadas por campo, en orden de preferencia. Se comparan normalizadas.
ALIAS_PROVEEDOR: dict[str, tuple[str, ...]] = {
    "id": ("id", "idproveedor", "proveedorid", "codigo", "codigoproveedor", "proveedor"),
    "razon_social": ("razonsocial", "razon", "nombre", "denominacion", "nombrecomercial"),
    "nif": ("nif", "cif", "nifcif", "cifnif", "documento", "identificacionfiscal"),
    "iban": ("iban", "cuenta", "cuentabancaria", "ccc", "ibanpago"),
    "ciudad": ("ciudad", "poblacion", "localidad", "provincia"),
    "condiciones_dias": (
        "condiciones",
        "condicionespago",
        "condicionesdepago",
        "dias",
        "diaspago",
        "diasdepago",
        "plazo",
        "plazopago",
        "plazodepago",
    ),
}
ALIAS_PEDIDO: dict[str, tuple[str, ...]] = {
    "pedido": ("pedido", "numeropedido", "numpedido", "npedido", "ordencompra", "orden", "po"),
    "proveedor_id": ("proveedorid", "idproveedor", "proveedor", "codigoproveedor"),
    "nif": ("nif", "cif", "nifcif", "cifnif", "documento"),
    "importe_total": ("importetotal", "importe", "total", "importepedido", "importeeuros"),
    "estado": ("estado", "situacion", "status"),
    "fecha_pedido": (
        "fechapedido",
        "fechadelpedido",
        "fecha",
        "fechaemision",
        "fechaorden",
        "fechaalta",
    ),
}

#: Sin estas columnas no se carga: el resto se puede suplir con avisos, esto no.
CLAVES_OBLIGATORIAS = {"Proveedores": ("id",), "Pedidos_2026": ("pedido", "importe_total")}
#: Su ausencia no impide cargar, pero hace escalar facturas: se avisa en mayúsculas.
COLUMNAS_CRITICAS = {"Proveedores": ("nif", "iban"), "Pedidos_2026": ("proveedor_id", "estado")}

_RE_NORMA = re.compile(r"norma|regla|politica|pagos")
_FILAS_CABECERA = 6  # hasta dónde se busca la fila de cabecera (por si hay un título encima)


class ErrorMaestro(KeyError):
    """El Excel no se puede leer como maestro. Hereda de KeyError por compatibilidad."""

    def __str__(self) -> str:
        return str(self.args[0]) if self.args else ""


def _txt(v: object) -> str:
    return "" if v is None else str(v).strip()


def _norm(v: object) -> str:
    return re.sub(r"[^a-z0-9]", "", sin_tildes(_txt(v)).lower())


def _localizar_hoja(wb: openpyxl.Workbook, canonico: str) -> str:
    """El nombre real de la hoja, aunque venga con otras mayúsculas, tildes o separadores."""
    por_nombre = {_norm(h): h for h in wb.sheetnames}
    for alias in ALIAS_HOJAS[canonico]:
        if alias in por_nombre:
            return por_nombre[alias]
    raise ErrorMaestro(
        f"el Excel no trae la hoja «{canonico}» (hay {wb.sheetnames}). "
        f"Nombres que valen: {', '.join(ALIAS_HOJAS[canonico])}"
    )


def _cabecera(
    filas: list[tuple[Any, ...]], alias: dict[str, tuple[str, ...]], hoja: str, avisos: list[str]
) -> tuple[int, dict[str, int]]:
    """Fila donde está la cabecera y columna de cada campo. Avisa de lo que sobra y lo que falta."""
    mejor: tuple[int, dict[str, int]] = (0, {})
    for i, fila in enumerate(filas[:_FILAS_CABECERA]):
        cols: dict[str, int] = {}
        usadas: set[int] = set()
        normalizadas = [_norm(c) for c in fila or ()]
        for campo, nombres in alias.items():
            for nombre in nombres:
                libres = [j for j, n in enumerate(normalizadas) if n == nombre and j not in usadas]
                if libres:
                    cols[campo] = libres[0]
                    usadas.add(libres[0])
                    break
        if len(cols) > len(mejor[1]):
            mejor = (i, cols)
    idx, cols = mejor
    if not cols:
        raise ErrorMaestro(
            f"{hoja}: no reconozco ninguna cabecera en las primeras {_FILAS_CABECERA} filas "
            f"(esperaba alguna de {', '.join(next(iter(alias.values())))}). "
            "Mira si la hoja tiene títulos arriba o si las columnas se llaman de otra forma."
        )
    if idx:
        avisos.append(f"{hoja}: la cabecera está en la fila {idx + 1}, no en la primera")
    faltan = [c for c in alias if c not in cols]
    for campo in faltan:
        critica = campo in COLUMNAS_CRITICAS.get(hoja, ()) or campo in CLAVES_OBLIGATORIAS.get(
            hoja, ()
        )
        avisos.append(
            f"{hoja}: FALTA la columna «{campo}»"
            + (" y sin ella las facturas afectadas escalarán" if critica else " (queda vacía)")
        )
    faltan_obligatorias = [c for c in CLAVES_OBLIGATORIAS.get(hoja, ()) if c not in cols]
    if faltan_obligatorias:
        raise ErrorMaestro(
            f"{hoja}: sin la columna «{faltan_obligatorias[0]}» no se puede cargar el maestro. "
            f"Cabeceras leídas en la fila {idx + 1}: {[_txt(c) for c in filas[idx] if c is not None]}"
        )
    conocidas = set(cols.values())
    desconocidas = [_txt(c) for j, c in enumerate(filas[idx]) if j not in conocidas and _norm(c)]
    if desconocidas:
        avisos.append(f"{hoja}: columnas que no conozco y no uso: {desconocidas}")
    return idx, cols


def _celda(fila: tuple[Any, ...], cols: dict[str, int], campo: str) -> Any:
    j = cols.get(campo)
    return fila[j] if j is not None and j < len(fila) else None


def _datos(filas: list[tuple[Any, ...]], idx: int) -> list[tuple[int, tuple[Any, ...]]]:
    """Filas de datos con su número real de fila en la hoja, sin las vacías."""
    return [
        (n, f)
        for n, f in enumerate(filas[idx + 1 :], start=idx + 2)
        if f and any(_txt(c) for c in f)
    ]


def cargar_maestro(ruta: Path | str) -> MasterSnapshot:
    with closing(openpyxl.load_workbook(ruta, data_only=True, read_only=True)) as wb:
        avisos: list[str] = []
        proveedores: dict[str, Proveedor] = {}
        pedidos: dict[str, Pedido] = {}

        hoja_prov = _localizar_hoja(wb, "Proveedores")
        if hoja_prov != "Proveedores":
            avisos.append(f"Proveedores: la hoja se llama «{hoja_prov}»; la leo igual")
        filas = list(wb[hoja_prov].iter_rows(values_only=True))
        idx, cols = _cabecera(filas, ALIAS_PROVEEDOR, "Proveedores", avisos)
        for n, fila in _datos(filas, idx):
            crudo_id = _celda(fila, cols, "id")
            if not _txt(crudo_id):
                avisos.append(
                    f"Proveedores: fila {n} sin ID, se ignora ({[_txt(c) for c in fila]})"
                )
                continue
            pid = _txt(crudo_id).upper()
            razon = _celda(fila, cols, "razon_social")
            m = re.search(r"\d+", _txt(_celda(fila, cols, "condiciones_dias")))
            p = Proveedor(
                id=pid,
                razon_social=_txt(razon),
                nif=normalizar_nif(_txt(_celda(fila, cols, "nif"))),
                iban=normalizar_iban(_txt(_celda(fila, cols, "iban"))),
                ciudad=_txt(_celda(fila, cols, "ciudad")) or None,
                condiciones_dias=int(m.group()) if m else None,
            )
            if str(razon) != _txt(razon):
                avisos.append(
                    f"Proveedores: {pid} tiene espacios sobrantes en la razón social ({razon!r})"
                )
            if pid in proveedores:
                igual = proveedores[pid] == p
                avisos.append(
                    f"Proveedores: {pid} duplicado "
                    + (
                        "(filas idénticas; se ignora la segunda)"
                        if igual
                        else f"CON DATOS DISTINTOS: {proveedores[pid]} vs {p}"
                    )
                )
                continue
            proveedores[pid] = p

        hoja_ped = _localizar_hoja(wb, "Pedidos_2026")
        if hoja_ped != "Pedidos_2026":
            avisos.append(f"Pedidos_2026: la hoja se llama «{hoja_ped}»; la leo igual")
        filas = list(wb[hoja_ped].iter_rows(values_only=True))
        idx, cols = _cabecera(filas, ALIAS_PEDIDO, "Pedidos_2026", avisos)
        for n, fila in _datos(filas, idx):
            crudo_num = _celda(fila, cols, "pedido")
            if not _txt(crudo_num):
                avisos.append(
                    f"Pedidos_2026: fila {n} sin número de pedido, se ignora ({[_txt(c) for c in fila]})"
                )
                continue
            num = normalizar_pedido(_txt(crudo_num))
            crudo_importe = _celda(fila, cols, "importe_total")
            importe = parse_importe_es(crudo_importe)
            if importe is None:
                avisos.append(f"Pedidos_2026: {num} sin importe legible ({crudo_importe!r})")
                continue
            pedido = Pedido(
                pedido=num,
                proveedor_id=_txt(_celda(fila, cols, "proveedor_id")).upper(),
                nif=normalizar_nif(_txt(_celda(fila, cols, "nif"))),
                importe_total=importe.quantize(CENT),
                estado=_txt(_celda(fila, cols, "estado")).upper(),
                fecha_pedido=parse_fecha_es(_txt(_celda(fila, cols, "fecha_pedido"))),
            )
            if num in pedidos:
                igual = pedidos[num] == pedido
                avisos.append(
                    f"Pedidos_2026: {num} duplicado "
                    + (
                        "(filas idénticas; se ignora la segunda)"
                        if igual
                        else f"CON DATOS DISTINTOS (se conserva la primera): {pedidos[num]} vs {pedido}"
                    )
                )
                continue
            pedidos[num] = pedido
            prov = proveedores.get(pedido.proveedor_id)
            if not pedidos[num].nif:
                avisos.append(
                    f"Pedidos_2026: {num} sin NIF en el Excel ({pedidos[num].proveedor_id} tiene {prov.nif if prov else '?'})"
                )
            elif prov and prov.nif != pedidos[num].nif:
                avisos.append(
                    f"Pedidos_2026: {num} lleva NIF {pedidos[num].nif} pero {prov.id} tiene {prov.nif}"
                )
            if prov is None:
                avisos.append(
                    f"Pedidos_2026: {num} referencia un proveedor que no está en el maestro ({pedidos[num].proveedor_id})"
                )

        for hoja in HOJAS_NOTAS:
            if hoja in wb.sheetnames:
                for fila in wb[hoja].iter_rows(values_only=True):
                    texto = " | ".join(_txt(c) for c in fila if c is not None)
                    if texto:
                        avisos.append(f"{hoja}: {texto}")
        leidas = {hoja_prov, hoja_ped, HOJA_NORMA_VIGENTE, *HOJAS_NOTAS}
        normas = hojas_norma(wb)
        for hoja in normas:
            if hoja not in leidas:
                avisos.append(
                    f"REGLA NUEVA?: la hoja «{hoja}» parece una norma y nadie la lee. "
                    f"Léela con `albertitos.sources.excel.leer_norma(xlsx, {hoja!r})` "
                    "antes de decidir nada: puede ser la regla del sábado."
                )
        ignoradas = sorted(set(wb.sheetnames) - leidas - set(normas))
        if ignoradas:
            avisos.append(f"hojas ignoradas: {ignoradas}")

        version = hash_canonico(
            {
                "p": {k: v.model_dump(mode="json") for k, v in proveedores.items()},
                "o": {k: v.model_dump(mode="json") for k, v in pedidos.items()},
            }
        )[:12]
        return MasterSnapshot(
            version=version,
            origen=str(ruta),
            proveedores=proveedores,
            pedidos=pedidos,
            avisos_calidad=avisos,
            cargado_en=datetime.now(UTC),
        )


def hojas_norma(libro: openpyxl.Workbook | Path | str) -> list[str]:
    """Hojas cuyo nombre suena a norma de pagos. La regla nueva del sábado puede llegar así."""
    if isinstance(libro, (str, Path)):
        with closing(openpyxl.load_workbook(libro, data_only=True, read_only=True)) as wb:
            return hojas_norma(wb)
    return [h for h in libro.sheetnames if _RE_NORMA.search(_norm(h))]


def leer_norma(ruta: Path | str, hoja: str = HOJA_NORMA_VIGENTE) -> list[str]:
    """Texto literal de la norma, para docs y para comparar con la v4 del sábado."""
    with closing(openpyxl.load_workbook(ruta, data_only=True, read_only=True)) as wb:
        por_nombre = {_norm(h): h for h in wb.sheetnames}
        real = por_nombre.get(_norm(hoja))
        if real is None:
            raise ErrorMaestro(
                f"el Excel no trae la hoja «{hoja}». Hojas que suenan a norma: {hojas_norma(wb)}"
            )
        return [_txt(f[0]) for f in wb[real].iter_rows(values_only=True) if f and f[0]]
