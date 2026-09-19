import importlib.util
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from albertitos.sources.excel import ErrorMaestro, cargar_maestro, hojas_norma, leer_norma

MAESTRO_XLSX = "FINAL_v7_DEFINITIVO_ahorasi.xlsx"
VERSION_REAL = "80911e429c6c"  # la que está en la BD y en la entrega publicada
_GENERAR = Path(__file__).parent.parent / "data/fixtures/maestro_cambiado/generar.py"


def _generador():
    """El generador de fixtures vive en data/fixtures/, que no es un paquete importable."""
    spec = importlib.util.spec_from_file_location("maestro_cambiado_generar", _GENERAR)
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return modulo


GEN = _generador()


@pytest.fixture(scope="session")
def variantes(caja, tmp_path_factory):
    """Las copias del Excel real con otra forma o otro contenido, generadas una vez."""
    destino = tmp_path_factory.mktemp("maestro_cambiado")
    nombres = ["original", *GEN.VARIANTES]
    return {n: GEN.generar(n, destino, caja / MAESTRO_XLSX) for n in nombres}


def test_maestro_real(caja):
    m = cargar_maestro(caja / "FINAL_v7_DEFINITIVO_ahorasi.xlsx")
    assert len(m.proveedores) == 11
    assert len(m.pedidos) == 516
    assert m.proveedores["P003"].razon_social == "Ofimática Cieza S.L."  # sin espacios sobrantes
    assert m.proveedores["P001"].iban == "ES2100491500051234567890"
    assert m.proveedores["P001"].condiciones_dias == 60
    assert m.pedidos["PO-2026-0001"].importe_total == Decimal("9221.75")
    assert "Proveedores: P007 duplicado (filas idénticas; se ignora la segunda)" in m.avisos_calidad
    assert any("P003 tiene espacios sobrantes" in a for a in m.avisos_calidad)
    assert "pendiente_revisar: PO-2026-0007" in m.avisos_calidad
    assert "pendiente_revisar: PO-2026-0141" in m.avisos_calidad
    assert {n for n, p in m.pedidos.items() if not p.nif} == {
        f"PO-2026-{n:04d}" for n in range(538, 558)
    }
    assert any("NUNCA pagar sin cruzar" in a for a in m.avisos_calidad)
    assert (
        sum(1 for a in m.avisos_calidad if "sin NIF en el Excel" in a) == 20
    )  # PO-2026-0538..0557
    assert m.proveedor_por_nif("b-46102331").id == "P001"
    assert len(m.version) == 12


def test_norma_v3_tiene_seis_reglas(caja):
    lineas = leer_norma(caja / "FINAL_v7_DEFINITIVO_ahorasi.xlsx")
    assert sum(1 for x in lineas if x[:2] in {f"{i}." for i in range(1, 7)}) == 6


PROVEEDOR = ["P001", "Proveedor uno", "B46102331", "ES2100491500051234567890", "Murcia", "60 dias"]
PEDIDO = ["PO-2026-0001", "P001", "B46102331", "1.234,56", "ABIERTO", "18/09/2026"]


def _crear_excel(tmp_path, *, proveedores=None, pedidos=None):
    ruta = tmp_path / "maestro.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Proveedores"
    wb.active.append(["id", "razon", "nif", "iban", "ciudad", "condiciones"])
    for fila in proveedores if proveedores is not None else [PROVEEDOR]:
        wb.active.append(fila)
    hoja = wb.create_sheet("Pedidos_2026")
    hoja.append(["pedido", "proveedor", "nif", "importe", "estado", "fecha"])
    for fila in pedidos if pedidos is not None else [PEDIDO]:
        hoja.append(fila)
    wb.create_sheet("Norma_Pagos_v3").append(["1. Una regla de prueba"])
    wb.save(ruta)
    wb.close()
    return ruta


def test_normaliza_maestro_sin_completar_nif_ausente(tmp_path):
    proveedor = [
        " p001 ",
        " Proveedor uno ",
        " b-46102331 ",
        " es21 0049 1500 0512 3456 7890 ",
        " Murcia ",
        " 60 dias ",
    ]
    pedido = [" po-2026- 0001 ", " p001 ", None, " 1.234,56 EUR ", " abierto ", date(2026, 9, 18)]
    m = cargar_maestro(_crear_excel(tmp_path, proveedores=[proveedor], pedidos=[pedido]))
    assert m.proveedores["P001"].nif == "B46102331"
    assert m.proveedores["P001"].iban == "ES2100491500051234567890"
    assert m.proveedores["P001"].razon_social == "Proveedor uno"
    assert m.proveedores["P001"].condiciones_dias == 60
    p = m.pedidos["PO-2026-0001"]
    assert p.proveedor_id == "P001"
    assert p.nif == ""
    assert p.importe_total == Decimal("1234.56")
    assert p.fecha_pedido == date(2026, 9, 18)
    assert p.estado == "ABIERTO"
    assert any("sin NIF" in a and "B46102331" in a for a in m.avisos_calidad)


@pytest.mark.parametrize("distinto", [False, True])
def test_proveedor_duplicado_conserva_primero_y_explica_conflicto(tmp_path, distinto):
    segundo = PROVEEDOR.copy()
    if distinto:
        segundo[2] = "B99999999"
    m = cargar_maestro(_crear_excel(tmp_path, proveedores=[PROVEEDOR, segundo]))
    assert len(m.proveedores) == 1
    assert m.proveedores["P001"].nif == "B46102331"
    aviso = next(a for a in m.avisos_calidad if "P001 duplicado" in a)
    assert ("CON DATOS DISTINTOS" if distinto else "filas idénticas") in aviso
    if distinto:
        assert "B46102331" in aviso and "B99999999" in aviso


@pytest.mark.parametrize("distinto", [False, True])
def test_pedido_duplicado_conserva_primero_y_explica_conflicto(tmp_path, distinto):
    segundo = PEDIDO.copy()
    if distinto:
        segundo[3] = "9.999,99"
    m = cargar_maestro(_crear_excel(tmp_path, pedidos=[PEDIDO, segundo]))
    assert len(m.pedidos) == 1
    assert m.pedidos["PO-2026-0001"].importe_total == Decimal("1234.56")
    aviso = next(a for a in m.avisos_calidad if "PO-2026-0001 duplicado" in a)
    assert ("CON DATOS DISTINTOS" if distinto else "filas idénticas") in aviso
    if distinto:
        assert "1234.56" in aviso and "9999.99" in aviso


def test_pedido_sin_proveedor_ni_nif_revela_ambas_carencias(tmp_path):
    pedido = PEDIDO.copy()
    pedido[1:3] = ["P999", None]
    m = cargar_maestro(_crear_excel(tmp_path, pedidos=[pedido]))
    assert any("sin NIF" in a for a in m.avisos_calidad)
    assert any("proveedor que no está" in a and "P999" in a for a in m.avisos_calidad)


@pytest.mark.parametrize("lector", [cargar_maestro, leer_norma])
@pytest.mark.parametrize("con_error", [False, True])
def test_cierra_excel_incluso_si_falta_hoja(tmp_path, monkeypatch, lector, con_error):
    ruta = _crear_excel(tmp_path)
    if con_error:
        wb = openpyxl.load_workbook(ruta)
        del wb["Proveedores" if lector is cargar_maestro else "Norma_Pagos_v3"]
        wb.save(ruta)
        wb.close()
    abrir = openpyxl.load_workbook
    libros = []

    def abrir_y_registrar(*args, **kwargs):
        wb = abrir(*args, **kwargs)
        libros.append(wb)
        return wb

    monkeypatch.setattr(openpyxl, "load_workbook", abrir_y_registrar)
    try:
        if con_error:
            with pytest.raises(KeyError):
                lector(ruta)
        else:
            lector(ruta)
        assert libros[0]._archive.fp is None
    finally:
        for wb in libros:
            wb.close()


# --- El Excel del sábado con otra forma: data/fixtures/maestro_cambiado/generar.py ---


@pytest.mark.parametrize("variante", sorted(GEN.FORMA))
def test_otra_forma_no_cambia_el_maestro(caja, variantes, variante):
    """Columnas, cabeceras, hojas y huecos son envoltorio: el maestro sale idéntico, dato a dato."""
    real = cargar_maestro(caja / MAESTRO_XLSX)
    m = cargar_maestro(variantes[variante])
    assert m.proveedores == real.proveedores
    assert m.pedidos == real.pedidos
    assert m.version == real.version == VERSION_REAL


@pytest.mark.parametrize("variante", sorted(GEN.CONTENIDO))
def test_otro_contenido_cambia_la_version(caja, variantes, variante):
    m = cargar_maestro(variantes[variante])
    assert m.version != VERSION_REAL


def test_columnas_reordenadas_no_mezclan_iban_con_ciudad(variantes):
    """El fallo silencioso que buscábamos: leyendo por posición, el IBAN acabaría en la ciudad."""
    m = cargar_maestro(variantes["columnas_reordenadas"])
    p = m.proveedores["P001"]
    assert p.iban == "ES2100491500051234567890"
    assert p.ciudad == "Valencia"
    assert p.nif == "B46102331"
    assert m.pedidos["PO-2026-0001"].importe_total == Decimal("9221.75")
    assert m.pedidos["PO-2026-0001"].fecha_pedido == date(2026, 1, 31)


def test_columna_desconocida_se_avisa_y_no_desplaza(variantes):
    m = cargar_maestro(variantes["columna_nueva"])
    assert m.proveedores["P001"].nif == "B46102331"
    assert any("Centro_Coste" in a and "no conozco" in a for a in m.avisos_calidad)
    assert any("Observaciones_2026" in a and "no conozco" in a for a in m.avisos_calidad)


def test_filas_vacias_se_saltan_y_la_fila_sin_id_se_explica(variantes):
    m = cargar_maestro(variantes["filas_vacias"])
    assert any("sin ID" in a and "se ignora" in a for a in m.avisos_calidad)
    assert any("sin número de pedido" in a for a in m.avisos_calidad)


def test_titulo_encima_encuentra_la_cabecera(variantes):
    m = cargar_maestro(variantes["titulo_encima"])
    assert any("la cabecera está en la fila 3" in a for a in m.avisos_calidad)
    assert not any("MAESTRO PROVEEDORES" in str(p.id) for p in m.proveedores.values())


def test_hoja_renombrada_se_lee_y_se_dice(variantes):
    m = cargar_maestro(variantes["hoja_renombrada"])
    assert any("la hoja se llama «PEDIDOS 2026»" in a for a in m.avisos_calidad)
    assert any("la hoja se llama «Maestro Proveedores»" in a for a in m.avisos_calidad)


def test_hoja_que_parece_norma_avisa_con_su_nombre(variantes):
    """La regla de las 18:00 puede llegar dentro del Excel: no puede pasar desapercibida."""
    m = cargar_maestro(variantes["hoja_norma_v4"])
    aviso = next(a for a in m.avisos_calidad if a.startswith("REGLA NUEVA?"))
    assert "Norma_Pagos_v4" in aviso
    assert "Norma_Pagos_v4" in hojas_norma(variantes["hoja_norma_v4"])
    assert not any(
        a.startswith("REGLA NUEVA?") for a in cargar_maestro(variantes["original"]).avisos_calidad
    )
    assert any(
        "firma del director financiero" in x
        for x in leer_norma(variantes["hoja_norma_v4"], "norma pagos V4")
    )


def test_iban_cambiado_solo_toca_a_su_proveedor(caja, variantes):
    real = cargar_maestro(caja / MAESTRO_XLSX)
    m = cargar_maestro(variantes["iban_cambiado"])
    assert m.proveedores["P001"].iban == "ES9121000418450200051332"
    assert {k: v for k, v in m.proveedores.items() if k != "P001"} == {
        k: v for k, v in real.proveedores.items() if k != "P001"
    }
    assert m.pedidos == real.pedidos


def test_pedido_anulado_y_proveedor_nuevo_llegan_al_maestro(variantes):
    assert cargar_maestro(variantes["pedido_anulado"]).pedidos["PO-2026-0001"].estado == "ANULADO"
    m = cargar_maestro(variantes["proveedor_nuevo"])
    assert m.proveedores["P099"].condiciones_dias == 30
    assert m.pedidos["PO-2026-0999"].proveedor_id == "P099"


def test_sin_la_columna_clave_dice_que_mirar(tmp_path):
    """Un maestro medio vacío escalaría las 500 en silencio: mejor un error con nombre."""
    ruta = _crear_excel(tmp_path)
    wb = openpyxl.load_workbook(ruta)
    wb["Pedidos_2026"]["A1"] = "referencia interna"
    wb.save(ruta)
    wb.close()
    with pytest.raises(ErrorMaestro) as e:
        cargar_maestro(ruta)
    assert "pedido" in str(e.value) and "referencia interna" in str(e.value)


def test_columna_critica_ausente_avisa_pero_carga(tmp_path):
    ruta = _crear_excel(tmp_path)
    wb = openpyxl.load_workbook(ruta)
    wb["Proveedores"]["D1"] = "no_es_un_iban"
    wb.save(ruta)
    wb.close()
    m = cargar_maestro(ruta)
    assert m.proveedores["P001"].iban == ""
    assert any("FALTA la columna «iban»" in a and "escalarán" in a for a in m.avisos_calidad)


def test_sin_hoja_de_pedidos_dice_los_nombres_que_valen(tmp_path):
    ruta = _crear_excel(tmp_path)
    wb = openpyxl.load_workbook(ruta)
    wb["Pedidos_2026"].title = "otra_cosa"
    wb.save(ruta)
    wb.close()
    with pytest.raises(ErrorMaestro) as e:
        cargar_maestro(ruta)
    assert "Pedidos_2026" in str(e.value) and "otra_cosa" in str(e.value)
